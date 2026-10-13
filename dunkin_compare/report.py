from __future__ import annotations

from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

from .core import Item, Store, normalized


def write_report(path: Path, collected: list[tuple[Store, list[Item], str]],
                 statuses: list[dict], collected_at: datetime, demo: bool = False,
                 capture_method: str = "automatic extraction"):
    workbook = Workbook()
    workbook.remove(workbook.active)
    timestamp = collected_at.isoformat(timespec="seconds")
    date = collected_at.date().isoformat()

    def sheet(name, headers, rows):
        ws = workbook.create_sheet(name)
        ws.append(headers)
        for row in rows:
            ws.append(row)
        # Treat all external text as text, including strings beginning with '='.
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="17365D")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        if ws.max_row > 1:
            table = Table(displayName=name.replace(" ", ""), ref=ws.dimensions)
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            ws.add_table(table)
        for column in ws.columns:
            letter = column[0].column_letter
            ws.column_dimensions[letter].width = min(65, max(14, max(len(str(c.value or "")) for c in column) + 2))
        for row in ws.iter_rows(min_row=2):
            for cell, header in zip(row, headers):
                if "percent" in header:
                    cell.number_format = "0.00%"
                elif "price" in header or "difference" in header:
                    cell.number_format = "0.00"
        return ws

    prices = []
    for store, items, source in collected:
        for item in items:
            prices.append([store.group, store.store_id, store.brand, store.address, store.market,
                           store.channel, item.category, item.name, item.variant,
                           float(item.price), item.currency, date, timestamp, source,
                           "SYNTHETIC DEMO" if demo else capture_method])
    sheet("Menu Prices", ["group", "store_id", "brand", "address", "market", "channel",
                          "category", "item", "variant", "price", "currency", "date",
                          "collected_at", "source_url", "data_type"], prices)

    # No fuzzy matching: sizes, currencies, channels, and market must agree.
    parents = defaultdict(list)
    competitors = defaultdict(list)
    for store, items, _ in collected:
        for item in items:
            key = (normalized(store.market), store.channel, normalized(store.brand),
                   normalized(item.category), normalized(item.name),
                   normalized(item.variant), item.currency)
            target = parents if store.group == "parent" else competitors
            target[key].append((store, item))
    comparison = []
    unmatched = []
    for key in sorted(parents.keys() | competitors.keys()):
        own = parents.get(key, [])
        other = competitors.get(key, [])
        if not own or not other or not key[5]:
            reason = "Size/variant unspecified; automatic comparison withheld" if not key[5] else "No exact matching item in the other group"
            for store, item in own + other:
                unmatched.append([store.group, store.store_id, item.category, item.name,
                                  item.variant, float(item.price), item.currency, reason])
            continue
        for parent, ours in own:
            for competitor, theirs in other:
                difference = theirs.price - ours.price
                percent = float(difference / ours.price) if ours.price else None
                comparison.append([parent.market, parent.channel, parent.store_id, parent.address,
                                   competitor.store_id, competitor.address, ours.category,
                                   ours.name, ours.variant, ours.currency, float(ours.price),
                                   float(theirs.price), float(difference), percent, date])
    sheet("Comparison", ["market", "channel", "parent_store_id", "parent_address",
                         "competitor_store_id", "competitor_address", "category", "item",
                         "variant", "currency", "parent_price", "competitor_price",
                         "difference_competitor_minus_parent", "difference_percent", "date"], comparison)
    sheet("Unmatched Items", ["group", "store_id", "category", "item", "variant", "price",
                              "currency", "reason"], unmatched)
    sheet("Store Status", ["group", "store_id", "address", "status", "item_count", "message", "collected_at"],
          [[s[k] for k in ("group", "store_id", "address", "status", "item_count", "message")] + [timestamp]
           for s in statuses])
    sheet("Read Me", ["topic", "details"], [
        ["Data", "SYNTHETIC DEMO: invented test prices; not Dunkin prices" if demo else "Observed prices from source URLs at collection time"],
        ["Collection method", capture_method],
        ["Location verification", "For user-reviewed browser capture, the user confirms the selected store and channel. The software does not independently verify them."],
        ["Comparison", "Positive difference means competitor is more expensive. Negative means competitor is cheaper."],
        ["Matching", "Exact normalized brand, category, item, variant, market, currency, and channel. No automatic comparison when variant is missing."],
        ["Coverage", "Check Store Status. Failed stores have no price rows. A collected page may omit options or categories; full-menu coverage needs live validation."],
        ["Prices", "Menu prices as displayed; taxes, fees, modifiers, and discounts are not inferred. Pickup and delivery are kept separate."],
        ["Time", "Workbook date/time is the collection run's completion time in America/New_York with UTC offset, not a price effective date. Interactive capture journals record each saved batch's time."],
    ])
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)

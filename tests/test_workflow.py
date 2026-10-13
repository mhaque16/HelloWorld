import json
import tempfile
import unittest
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

from dunkin_compare.core import Item, Store, load_stores, parse_jsonld, price_amount
from dunkin_compare.report import write_report
from tests.support import PARENT, COMPETITOR

ROOT = Path(__file__).resolve().parent.parent


class WorkflowTests(unittest.TestCase):
    def setUp(self):
        self.parent = PARENT
        self.competitor = COMPETITOR
        self.html = (ROOT / "tests/fixtures/parent.html").read_text()

    def test_parse_store_menu_and_variants(self):
        items = parse_jsonld(self.html, self.parent.address)
        self.assertEqual(len(items), 3)
        self.assertEqual(items[0], Item("Coffee", "Hot Coffee", "Small", Decimal("2.49"), "USD"))

    def test_wrong_store_is_rejected(self):
        with self.assertRaisesRegex(ValueError, "selected store"):
            parse_jsonld(self.html, self.competitor.address)

    def test_ambiguous_price_rejected(self):
        for value in ("From $2.49", "$2.49 - $3.49", "1,299", "NaN", "-1", "", "2.499"):
            with self.subTest(value=value), self.assertRaises(ValueError):
                price_amount(value)

    def test_missing_price_and_category_rejected(self):
        for html in (self.html.replace('"price": "2.49",', ''),
                     self.html.replace('"name": "Coffee",', '"name": "",')):
            with self.assertRaises(ValueError):
                parse_jsonld(html, self.parent.address)

    def test_conflicting_prices_rejected(self):
        doc = json.loads(self.html.split('<script type="application/ld+json">')[1].split('</script>')[0])
        items = doc['hasMenu']['hasMenuSection'][0]['hasMenuItem']
        item = json.loads(json.dumps(items[0]))
        item['offers'][0]['price'] = '9.99'
        items.append(item)
        with self.assertRaisesRegex(ValueError, "Conflicting"):
            parse_jsonld('<script type="application/ld+json">'+json.dumps(doc)+'</script>', self.parent.address)

    def test_report_values_and_safe_text(self):
        parents = parse_jsonld(self.html, self.parent.address)
        competitors = parse_jsonld((ROOT / 'tests/fixtures/competitor.html').read_text(), self.competitor.address)
        parents += [Item('Coffee', '=HYPERLINK("bad")', 'Small', Decimal('1'), 'USD')]
        statuses = [{"group": "parent", "store_id": "missing", "address": "unknown",
                     "status": "failed", "item_count": 0, "message": "No prices"}]
        with tempfile.TemporaryDirectory() as temp:
            output = Path(temp) / 'test.xlsx'
            write_report(output, [(self.parent, parents, 'fixture'), (self.competitor, competitors, 'fixture')],
                         statuses, datetime.now(ZoneInfo('America/New_York')), True)
            wb = load_workbook(output)
            self.assertEqual(wb['Menu Prices'].max_row, 8)
            rows = list(wb['Comparison'].values)
            headers = rows[0]
            hot = next(dict(zip(headers, r)) for r in rows[1:] if r[7] == 'Hot Coffee')
            self.assertAlmostEqual(hot['difference_competitor_minus_parent'], .30)
            self.assertAlmostEqual(hot['difference_percent'], .30/2.49)
            self.assertEqual(wb['Store Status']['D2'].value, 'failed')
            self.assertIn('SYNTHETIC DEMO', wb['Read Me']['B2'].value)
            self.assertEqual(wb['Menu Prices']['H5'].data_type, 's')
            wb.close()

    def test_different_size_channel_and_currency_not_compared(self):
        base = Item('Coffee', 'Coffee', 'Small', Decimal('2'), 'USD')
        variations = [Item('Coffee', 'Coffee', 'Large', Decimal('3'), 'USD'),
                      Item('Coffee', 'Coffee', 'Small', Decimal('3'), 'CAD'),
                      Item('Coffee', 'Coffee', '', Decimal('3'), 'USD')]
        with tempfile.TemporaryDirectory() as temp:
            for index, variation in enumerate(variations):
                output = Path(temp) / f'{index}.xlsx'
                write_report(output, [(self.parent, [base], ''), (self.competitor, [variation], '')],
                             [], datetime.now(ZoneInfo('America/New_York')))
                wb = load_workbook(output)
                self.assertEqual(wb['Comparison'].max_row, 1)
                self.assertEqual(wb['Unmatched Items'].max_row, 3)
                wb.close()
            delivery = Store(**{**self.competitor.__dict__, 'channel': 'delivery'})
            output = Path(temp) / 'channel.xlsx'
            write_report(output, [(self.parent, [base], ''), (delivery, [base], '')],
                         [], datetime.now(ZoneInfo('America/New_York')))
            wb = load_workbook(output)
            self.assertEqual(wb['Comparison'].max_row, 1)
            wb.close()


if __name__ == '__main__':
    unittest.main()

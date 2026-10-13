import tempfile
import unittest
import csv
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from dunkin_compare.__main__ import main
from dunkin_compare.core import parse_jsonld
from tests.support import PARENT, COMPETITOR

ROOT = Path(__file__).resolve().parent.parent


class InteractiveCliTests(unittest.TestCase):
    def test_partial_capture_and_skipped_store_export_with_nonzero_exit(self):
        def collect(store, browser, journal):
            if store.group == 'competitor':
                return [], '', 'skipped', 'No prices saved'
            items = parse_jsonld((ROOT / 'tests/fixtures/parent.html').read_text(), store.address)
            return items, 'https://example.com/menu', 'partial', 'Browser closed before finishing'

        with tempfile.TemporaryDirectory() as temp:
            output = Path(temp) / 'interactive.xlsx'
            for store in (PARENT, COMPETITOR):
                with (Path(temp) / f'{store.group}.csv').open('w', newline='') as f:
                    row = {k: v for k, v in store.__dict__.items() if k != 'group'}
                    writer = csv.DictWriter(f, fieldnames=list(row))
                    writer.writeheader()
                    writer.writerow(row)
            with patch('dunkin_compare.interactive.collect_interactive', side_effect=collect):
                result = main(['--interactive', '--output', str(output),
                               '--parent', str(Path(temp) / 'parent.csv'),
                               '--competitors', str(Path(temp) / 'competitor.csv'),
                               '--profiles', str(Path(temp) / 'not-needed.json')])
            self.assertEqual(result, 1)
            wb = load_workbook(output)
            self.assertEqual(wb['Menu Prices'].max_row, 4)
            self.assertEqual(wb['Menu Prices']['O2'].value, 'user-reviewed browser capture')
            self.assertEqual(wb['Store Status']['D2'].value, 'partial')
            self.assertEqual(wb['Store Status']['D3'].value, 'skipped')
            self.assertEqual(wb['Comparison'].max_row, 1)
            wb.close()


if __name__ == '__main__':
    unittest.main()
